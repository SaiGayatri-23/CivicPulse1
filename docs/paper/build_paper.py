"""Builds the CivicPulse IEEE-format product paper as LaTeX (for Overleaf) and Word (.docx).

Run:  python docs/paper/build_paper.py
Outputs (next to this file): civicpulse_ieee.tex, CivicPulse_IEEE.docx, fig_architecture.png, fig_lifecycle.png
"""
import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

OUT = Path(__file__).parent
FONT = "C:/Windows/Fonts/times.ttf"
FONT_B = "C:/Windows/Fonts/timesbd.ttf"

TITLE = "CivicPulse: A Case Study in Building an AI-Assisted, Accountability-First Civic Engagement Platform for Hyderabad"

# Fill these in before submission.
AUTHORS = [
    ("1st Author Name", "Dept. of Computer Science and Engineering", "Name of Institution", "Hyderabad, India", "email@example.com"),
    ("2nd Author Name", "Dept. of Computer Science and Engineering", "Name of Institution", "Hyderabad, India", "email@example.com"),
    ("3rd Author Name", "Dept. of Computer Science and Engineering", "Name of Institution", "Hyderabad, India", "email@example.com"),
]

ABSTRACT = (
    "Citizens in fast-growing Indian cities routinely encounter potholes, overflowing drains, broken street lights and "
    "uncollected garbage, yet most complaint channels offer little visibility into what happens after a complaint is filed. "
    "This paper presents CivicPulse, a web and mobile-ready civic engagement platform for Hyderabad that treats every report "
    "as an accountable ticket rather than a one-way message. Citizens can report a problem with mandatory photographic "
    "evidence, a map location and optional video, with or without creating an account. Each report moves through an "
    "auditable lifecycle in which a fix can only be marked resolved with an after photograph, and the original reporter "
    "can confirm or dispute the fix. Three generative AI features built on a free-tier large language model reduce effort "
    "on both sides: photo-based auto-fill of the report form, semantic duplicate detection against nearby open reports, "
    "and department routing suggestions for administrators. The platform also provides a press-and-hold emergency button, "
    "missing child and city alerts, events, utility helplines, bus stop lookup, petitions and community contributions from "
    "verified organisations. Security is enforced in the database through row-level security policies, column-level "
    "grants and encrypted storage of personal data. The implementation uses React with TypeScript on the client and a "
    "managed PostgreSQL backend. Presented as a case study, the paper traces the problem, the design objectives, the "
    "development journey from an early prototype to a database-centred rebuild, the complete end-to-end workflow "
    "from the moment a citizen notices a problem to the moment the fix is confirmed, the implementation, the challenges "
    "faced, an assessment of how far each objective was met, lessons learned, limitations and future scope."
)

INDEX_TERMS = ("civic technology, e-governance, grievance redressal, crowdsourced reporting, large language models, "
               "row-level security, progressive web application, smart cities")

# Content blocks: ("h1", text) ("h2", text) ("p", text) ("bullets", [items]) ("table", caption, header, rows) ("fig", file, caption)
BODY = [
    ("h1", "Introduction"),
    ("p", "Urban local bodies such as the Greater Hyderabad Municipal Corporation (GHMC) receive a steady stream of "
          "grievances about roads, sanitation, street lighting, water supply and public spaces. Existing channels, "
          "including call centres, social media and general-purpose complaint portals, share three weaknesses. First, the "
          "citizen rarely learns what happened after filing a complaint. Second, the same problem is reported many times "
          "by different people, which wastes staff effort. Third, there is no reliable proof that a problem was actually "
          "fixed, so tickets can be closed on paper while the problem remains on the street."),
    ("p", "CivicPulse addresses these gaps with an accountability-first design. Every report receives a public reference "
          "number, a visible status pipeline and a promised fix date. A report cannot be marked resolved without an after "
          "photograph, and the reporter has seven days to confirm or dispute the fix. Disputed fixes reopen the ticket and "
          "are flagged across the public feed. Generative AI is used only where it removes friction and never takes an "
          "action on its own: every AI output is a suggestion that a person accepts or discards."),
    ("p", "The main contributions of this work are: (a) an end-to-end reporting and resolution workflow with enforced "
          "photographic evidence and reporter verification; (b) three practical, cost-bounded AI assistants for citizens "
          "and administrators; (c) a privacy-preserving design in which personal data is encrypted at rest and access "
          "rules are enforced by the database rather than the client; and (d) a single platform that combines grievance "
          "reporting with public safety and community services for one city."),

    ("h1", "Case Background and Problem Statement"),
    ("h2", "Context"),
    ("p", "Hyderabad is one of India's fastest-growing metropolitan areas, and its municipal corporation, GHMC, is "
          "responsible for roads, drainage, street lighting, waste collection and public spaces across a large and "
          "expanding area. Residents already report problems through a call centre, a WhatsApp number, a web portal and "
          "social media. In practice, however, a resident who reports a broken street light has no simple way to see who "
          "is handling it, whether anyone else has already reported it, when it will be fixed, or whether the ticket was "
          "closed because the light was actually repaired."),
    ("h2", "Motivation"),
    ("p", "The project began from an everyday frustration shared by many residents: problems on neighbourhood streets, "
          "such as a pothole or a street light that stays broken for weeks, are reported but not visibly followed up, and "
          "there is no way to tell whether a closed complaint meant a real repair. The authors wanted to find out whether a small team using only free tools could build a platform that makes civic "
          "follow-through visible, and whether modern AI could make reporting easier without taking decisions away from "
          "people. A second motivation was to gather scattered civic information (emergency numbers, alerts, events, "
          "utilities and transport) into one place that residents would open every day, not only when something breaks."),
    ("h2", "Problem Statement"),
    ("p", "The case addressed in this project is therefore not the absence of a complaint channel, but the absence of "
          "visibility, proof and follow-through after a complaint is made. The project set out to design and build a "
          "single citizen-facing platform that makes every civic problem traceable from report to verified fix, reduces "
          "duplicate and low-quality reports, protects the personal data of the people who report, and also gives "
          "residents everyday civic information (safety alerts, utilities, transport and community activity) in one place."),
    ("h2", "Objectives"),
    ("bullets", [
        "O1 Transparency: every report has a public reference, a visible status timeline and a promised fix date.",
        "O2 Proof of resolution: a report cannot be closed as fixed without an after photograph, and the reporter can dispute the fix.",
        "O3 Low effort and inclusion: anyone can report in a few steps, without an account, in their own language, even when offline.",
        "O4 Less noise: duplicate reports are caught before and at submission, and AI assists reporting and triage.",
        "O5 Privacy and safety by design: personal data is encrypted and access rules are enforced by the database.",
        "O6 Sustainability: the platform runs on free service tiers and is funded by non-tracking sponsorship.",
    ]),
    ("h2", "Scope"),
    ("p", "The case covers one city (Hyderabad), web and installable mobile use through a progressive web application, "
          "and four kinds of users: citizens or guests, verified organisations, administrators and advertisers. CivicPulse "
          "is a community platform, not an official government channel, so it hands reports to GHMC through GHMC's own "
          "public channels rather than an internal integration."),

    ("h1", "Related Work"),
    ("p", "Research on citizen participation frames reporting platforms as a form of coproduction, in which residents act "
          "as partners in delivering public services rather than as customers; Linders describes this shift from "
          "e-government to \"we-government\" and classifies citizen sourcing as one of its main forms [13]. Goodchild "
          "showed that citizens can act as distributed sensors that contribute useful geographic information [14], which "
          "is the premise of location-tagged reporting. Nam and Pardo argue that smart city initiatives depend on people "
          "and institutions as much as on technology [15], which motivates the emphasis in CivicPulse on accountability "
          "and trust rather than technology alone."),
    ("p", "Open311 defines an open standard, GeoReport v2, for submitting and querying non-emergency service requests "
          "[1]. FixMyStreet, developed by mySociety in the United Kingdom, popularised map-based reporting that forwards "
          "reports to the responsible council [2]; an early evaluation found that it let citizens report, discuss and track the resolution of local problems, shifting some influence toward residents [16]. SeeClickFix offers a similar commercial service in North America "
          "[3]. In India, the Swachhata-MoHUA application focuses on sanitation complaints for urban local bodies [4]. "
          "These systems already give reporters some voice after closure: FixMyStreet emails the reporter four weeks "
          "later to ask whether the problem was fixed [17], and SeeClickFix lets a reporter reopen a closed issue for "
          "about seven days [18]. However, a ticket can still be closed without any photographic proof of the repair, and "
          "none of them uses language models to reduce reporting or triage effort. CivicPulse builds on these ideas and "
          "adds a database-enforced after photograph, a structured confirm-or-dispute step, semantic duplicate detection "
          "and database-enforced privacy controls. The Evaluation section compares the platforms directly."),

    ("h1", "Requirements and Stakeholders"),
    ("h2", "Stakeholders"),
    ("bullets", [
        "Citizens and guests who report problems, follow progress and confirm fixes.",
        "Verified organisations (NGOs, community groups, schools, businesses and government bodies) who host events and share community contributions.",
        "Administrators and department staff who triage, assign, resolve and moderate.",
        "Advertisers who run budget-limited advertising campaigns that help fund the service.",
    ]),
    ("h2", "Functional Requirements"),
    ("bullets", [
        "Report a problem with at least one photograph, a map pin, category, title and description, without requiring an account.",
        "Show a transparent status timeline and allow the reporter to accept or dispute a fix.",
        "Detect likely duplicate reports before and at submission time.",
        "Provide administrators with a ticket queue, claim locking, department assignment, merging of duplicates, moderation, verification review and an audit log.",
        "Provide public safety features: an emergency button, missing child alerts and city-wide service alerts.",
    ]),
    ("h2", "Non-Functional Requirements"),
    ("bullets", [
        "Privacy: mobile numbers, addresses and identity document digits must never be publicly readable.",
        "Scalability: administrative lists must remain usable with several hundred thousand users.",
        "Accessibility and reach: a choice of 23 Indian languages (20 translated so far, with right-to-left layout for Urdu, Kashmiri and Sindhi), dark, light and device themes, and phone, tablet and desktop layouts.",
        "Resilience: reports made offline must be queued on the device and sent automatically later.",
        "Cost: operate on free tiers of the database, AI and email services during the pilot phase.",
    ]),

    ("h1", "Development Methodology"),
    ("p", "CivicPulse was built iteratively in two phases, each ending with a working version that was tested before "
          "the next round of features began."),
    ("h2", "Phase 1: Prototype"),
    ("p", "The first version, built in June 2026, was a static web page that grew into a MERN application (MongoDB "
          "Atlas, Express, React and Node.js) with its own login and REST endpoints. It demonstrated the idea of "
          "map-based reporting and was made responsive for phones. Building it exposed the limits of that design: "
          "authentication and access checks lived in hand-written server code and needed repeated fixes, caching data in "
          "the browser overflowed the storage quota, demonstration administrator details were visible in the page, and "
          "configuration secrets had been committed to the repository. Every new feature added more server code that "
          "had to repeat the same security checks."),
    ("h2", "Phase 2: Rebuild on a Database-Centred Architecture"),
    ("p", "In September 2026 the prototype was retired and the platform was rebuilt from scratch with React, TypeScript "
          "and Tailwind CSS on the client and Supabase on the backend. The key design decision was to move every rule "
          "into PostgreSQL itself, as row-level security policies, triggers and functions, so that no custom server "
          "had to be written or secured. Configuration secrets that the prototype had committed to its "
          "repository were removed, and all secrets used by the new backend are kept in the database vault. The schema was then grown through 53 small, versioned migrations, each of which could be reviewed "
          "and applied on its own."),
    ("h2", "Process and Tools"),
    ("bullets", [
        "Feature cycle: each feature was specified as screens and rules, implemented as a database migration plus client pages, then tested with impersonated users and in the browser at phone, tablet and desktop widths before moving on.",
        "Version control and integration: Git with a continuous integration workflow that runs the unit tests, type-checks and builds the client, and audits production dependencies for known vulnerabilities on every push.",
        "Free-tier constraint as a design input: every service (database, AI, email, maps, CAPTCHA) was chosen for a usable free tier, and AI and email usage limits were enforced in the database.",
        "AI-assisted development: an AI coding assistant was used for pair programming, code review and drafting of migrations, with every change reviewed, tested and accepted by the authors.",
    ]),

    ("h1", "End-to-End Workflow"),
    ("p", "This section describes what the application does from a user's point of view, before the implementation is "
          "discussed. The core journey is the life of a single report, followed by the other journeys the platform supports."),
    ("h2", "The Core Journey: From Problem to Verified Fix"),
    ("bullets", [
        "Notice: a resident sees a large pothole on their street and opens CivicPulse on their phone. No sign-in is needed; a guest session is created in the background.",
        "Describe: they photograph the pothole. The location is read from the photo, and AI auto-fill suggests the category, a title and a description, which the resident can edit. Faces and number plates can be blurred before upload.",
        "Check for duplicates: the form lists open reports of the same category within 250 metres and, on request, the AI compares the meaning of the draft with them. If the pothole is already reported, the resident backs that report instead of filing a new one.",
        "Submit: the report receives a public reference number and appears on the feed and map. If the phone is offline, it is queued and sent automatically later. The resident is offered a one-tap hand-off to GHMC's WhatsApp and grievance portal.",
        "Triage: an administrator claims the ticket, accepts or overrides the AI department suggestion, and sets a promised fix date. Each step appears on the public timeline, and the resident is notified in the app, by web push or by email.",
        "Community: neighbours back and follow the report, add their own photographs and comment; staff replies are marked as official.",
        "Resolve: the department fixes the pothole and the administrator uploads an after photograph; without it, the database refuses to mark the report resolved.",
        "Verify: the resident sees the before and after photographs and has seven days to confirm or dispute. A dispute with a reason reopens the ticket and marks it as a disputed fix everywhere it appears; silence counts as confirmation.",
        "Account: confirmed fixes feed the public department scorecard (on-time rate, days to fix, confirmed share), and the resident earns points toward their tier.",
    ]),
    ("p", "This journey corresponds to the lifecycle in Fig. 2 and to the numbered steps on the architecture in Fig. 1."),
    ("h2", "Other Journeys"),
    ("p", "Table I lists what each type of user can do, and the scenarios below describe the most important journeys other than reporting."),
    ("table", "Principal Use Cases by Actor",
        ["Actor", "What they do in CivicPulse"],
        [
            ["Citizen or guest", "Report, back, follow and comment on problems; confirm or dispute fixes; watch areas; RSVP to events; sign petitions; use SOS, utilities and bus stops"],
            ["Verified individual", "All of the above, plus a verified badge and posting community contributions"],
            ["Verified organisation", "Host events, post community contributions, give official replies (government bodies)"],
            ["Administrator", "Triage, assign, resolve and merge reports; moderate; review verifications; issue child and city alerts; manage content, ads and users"],
            ["Advertiser", "Create, submit and track budget-limited campaigns"],
        ]),
    ("bullets", [
        "Emergency: a resident in danger presses and holds the SOS button for 1.2 seconds; the phone dials 112 and offers to send their location by SMS to up to three saved contacts.",
        "Missing child: an administrator publishes an alert backed by a police FIR; it appears at the top of every page for 72 hours, and any resident can submit a sighting that only administrators can read.",
        "Community life: a verified NGO creates a clean-up drive, residents RSVP with the number of adults and children, and afterwards the NGO posts photographs of the completed work.",
        "Everyday information: a resident checks a scheduled water cut in their area, finds the local police station number, pays a utility bill through the official link, or looks up the nearest bus stop.",
        "Oversight: at the end of the week, an administrator reviews the dashboard for overdue tickets, hotspots and citizen satisfaction, and exports it as a CSV file.",
    ]),

    ("h1", "System Architecture"),
    ("p", "CivicPulse follows a thin-client, thick-database architecture, shown in Fig. 1. The client is a React 19 "
          "single-page application written in TypeScript, styled with Tailwind CSS and built with Vite. It is installable "
          "as a progressive web application. A Workbox service worker precaches the application shell and serves "
          "photographs and map tiles cache-first and API reads network-first with a short timeout, so recently seen "
          "reports remain readable offline. Reports made without a connection, including their photographs and "
          "video, are stored in an IndexedDB outbox and sent automatically when the connection returns. Maps are "
          "rendered with Leaflet on OpenStreetMap tiles, place search and reverse geocoding use Nominatim [8], and "
          "location is read from photo metadata with the exifr library."),
    ("fig_wide", "fig_architecture.png", "System architecture of CivicPulse: users, client modules, Supabase platform services, database layer and external services."),
    ("p", "The backend is a managed PostgreSQL database provided by Supabase [6], which also supplies authentication, "
          "object storage and an automatically generated REST interface. Business rules live in the database as "
          "triggers and security-definer functions, so the same rules apply regardless of which client calls the API. "
          "The deployed schema contains 44 tables, 109 application functions, 53 triggers and 124 row-level security "
          "policies, created through 53 versioned migrations. Secrets such as the AI and email API keys are kept in an encrypted vault inside the "
          "database and are never shipped to the browser. Calls to the language model are made synchronously from "
          "the database through the HTTP extension, while outgoing email is sent asynchronously so that it can never "
          "block a user action. Push notifications are sent the same way, by a trigger on the notifications table. "
          "A single scheduled job runs nightly to delete data that has passed its retention period. "
          "Table II lists the technology stack."),
    ("table", "Technology Stack",
        ["Layer", "Technology"],
        [
            ["Client", "React 19, TypeScript, Tailwind CSS 4, Vite, React Router, TanStack Query, Lucide icons"],
            ["PWA and offline", "Workbox service worker, IndexedDB outbox, Web Push"],
            ["Maps and location", "Leaflet, React Leaflet, OpenStreetMap, Nominatim, Overpass API, exifr"],
            ["Backend", "Supabase: PostgreSQL, Auth (email, guest, optional Google), Storage, REST API"],
            ["Database extensions", "pgcrypto, Vault, pg_net, http, pg_cron, pg_trgm"],
            ["AI", "Google Gemini (free tier) via database HTTP calls"],
            ["Email", "Resend, sent asynchronously from the database"],
            ["Weather", "Open-Meteo weather and air quality APIs"],
            ["Bot defence", "Cloudflare Turnstile CAPTCHA"],
            ["Testing", "Vitest unit tests, TypeScript type checking"],
        ]),

    ("h1", "Core Modules"),
    ("h2", "Accounts and Identity"),
    ("p", "Users sign up with email and password, with email confirmation and password reset; Google sign-in can be "
          "enabled as an option. At sign-up a user chooses one of six account types: individual, community group, "
          "NGO, school or college, local business, or government body. They also give a mobile number and a "
          "structured address (house number, building, street, area, city and PIN code), which can be filled in from "
          "the device location. A guest who has been tracking reports on a device can convert that guest session into "
          "a full account and keep those reports. Verified accounts receive a badge, can upload a profile picture and "
          "can share community contributions; verified organisations can also host events."),
    ("h2", "Discovery and the Public Feed"),
    ("p", "The home screen shows active safety alerts, a city-at-a-glance panel with open, resolved and fix-rate "
          "figures and their week-on-week change, a personal panel of the user's own reports and unread updates, "
          "recently fixed reports with before and after photographs, upcoming events, and local weather and air "
          "quality from Open-Meteo. The report feed can be viewed as a grid, a list or a map. It can be filtered by "
          "category, status, time period, area and distance, searched by text or reference number, and sorted by "
          "latest, recently updated, most backed, nearest, needs attention or followed. Reports can be shared through "
          "the device share sheet or WhatsApp."),
    ("h2", "Issue Reporting"),
    ("p", "The report form collects a category, quick-detail chips, a title, a description, one to three photographs, an "
          "optional video of up to 30 seconds, and a map location. The location can be taken from the photograph's "
          "embedded GPS data, the device location or a tap on the map; hidden location metadata is stripped from photos "
          "before upload, and a built-in editor lets reporters blur faces and number plates. Reporters choose whether "
          "the report is public, anonymous (name hidden from the public but visible to administrators) or confidential "
          "(hidden from the public feed entirely). Guests can report without an account through an anonymous session, "
          "limited to three reports in 30 days and protected by a CAPTCHA [12]; their contact details are encrypted and "
          "stored in an administrator-only table. A priority score is computed from the category and hazard keywords. "
          "Reporters can edit or delete their own report while it is still pending, other residents can add their own "
          "photographs of the same problem, and a one-tap panel helps the reporter also file the complaint through "
          "GHMC's official WhatsApp and online grievance channels, since CivicPulse itself is not an official channel."),
    ("h2", "Accountable Resolution Lifecycle"),
    ("p", "Every report follows the lifecycle in Fig. 2. Administrators claim a ticket before editing it; a claim lasts "
          "four hours and prevents two staff members from working on the same report. Assigning a department and a "
          "promised fix date are recorded on a public timeline, and reports past their promised date are marked overdue. "
          "A report can only move to resolved when a proof photograph is attached, a rule enforced by a database trigger "
          "rather than the user interface. The reporter then has seven days to confirm or dispute the fix. A dispute "
          "requires a short explanation, returns the ticket to in progress, increments a reopen counter and displays a "
          "disputed fix badge on the report wherever it appears. If the reporter does not respond within seven days, "
          "the fix is treated as confirmed; this is evaluated when the data is read rather than by a scheduled job. "
          "Reports that cannot be fixed must be closed with one of five stated reasons, such as private land or "
          "another agency. A public scorecard ranks departments by open, overdue and resolved reports, average days "
          "to fix, on-time rate and the share of fixes confirmed by reporters."),
    ("fig", "fig_lifecycle.png", "Report lifecycle with enforced proof and reporter verification."),
    ("h2", "Duplicate Handling"),
    ("p", "Duplicates are handled in three layers. While the reporter is still filling in the form, open reports of the "
          "same category within 250 metres are listed so that the reporter can back an existing report instead. At "
          "submission, a database trigger refuses a report that repeats an open report within 50 metres, unless the "
          "reporter explicitly confirms that theirs is a different problem; a guest who reports the same problem twice "
          "is matched through hashed contact details. Finally, administrators can merge a duplicate into the original, "
          "which moves its backers and followers and notifies the reporter."),
    ("h2", "Engagement and Communication"),
    ("p", "Citizens can back, follow and comment on reports, with replies nested one level deep. Official replies from "
          "staff and verified government bodies are visually labelled, and staff can insert standard replies. Comments "
          "can be edited for 15 minutes and are limited to three per report per person to prevent spam. Updates reach "
          "users through an in-app notification centre, optional web push notifications and optional email alerts, "
          "which also allow guests to follow a report from any device without an account. Users can also watch up to "
          "five areas and be notified of new public reports within a chosen distance of each. A help and feedback "
          "section provides message threads between users and administrators."),
    ("h2", "Participation and Recognition"),
    ("p", "Useful actions earn points recorded in a ledger: for example, reporting an issue, having a report resolved, "
          "confirming a fix, commenting and completing verification. Points place users in tiers shown on their "
          "profile, time-bound missions award extra points for specific tasks, and the missions page lists the most "
          "active contributors of the month."),

    ("h1", "AI-Assisted Features"),
    ("p", "All AI features use a free-tier Gemini model [7] through structured JSON output, so every response is parsed "
          "against a fixed schema instead of free text. The model name is configurable from the vault, which allowed the "
          "system to move to a newer model when an older one was retired without a code change. Each feature has "
          "per-user and global daily limits recorded in a usage log, and no AI output is ever applied without a person "
          "accepting it. Table III summarises the three features."),
    ("table", "AI Features",
        ["Feature", "User", "Input and output"],
        [
            ["Photo auto-fill", "Citizen", "Photo in; category, title and description out"],
            ["Duplicate check", "Citizen", "Draft text and nearby reports in; best match and reason out"],
            ["Department routing", "Admin", "Report text in; department, confidence and reason out"],
        ]),
    ("p", "Photo auto-fill downsizes the first photograph on the device to 768 pixels before sending it, keeping requests "
          "small and fast. The duplicate check complements the geographic rules described above by comparing meaning "
          "rather than distance: in testing, for a draft describing a street light left on during the day, it correctly "
          "matched an existing report titled \"Light on during the day\" and rejected a nearby report about street "
          "lights being off, which is in the same category and area but is the opposite problem. Department routing "
          "results are cached against a hash of the report text so that repeated clicks do not consume quota."),

    ("h1", "Public Safety and Community Services"),
    ("h2", "Safety"),
    ("bullets", [
        "Emergency button: a press-and-hold control (1.2 seconds, to avoid accidental activation) on every page that places a normal call to 112, shows the user's location and helps share it by call or SMS with up to three saved emergency contacts. It never claims that help has been dispatched.",
        "Missing child alerts: issued by administrators only with a police FIR reference, which is never shown publicly. They appear at the top of every page with one-tap calling of 112 and the police station, expire after 72 hours unless renewed, and have identifying details removed when closed. Anyone can report a sighting, and only administrators can read sightings.",
        "City alerts: area-specific notices such as scheduled water supply cuts, with an expiry time.",
    ]),
    ("h2", "Community Services"),
    ("bullets", [
        "Utilities: national helplines, local police stations matched to the user's area, and official bill payment links.",
        "Bus stops and routes: nearby stops from OpenStreetMap through the Overpass API, routes curated by administrators with next-bus estimates, and a link to the transit operator's live tracking app.",
        "Events: civic, health, animal welfare and children's events with capacity-limited RSVP that records adults and children, a family-friendly filter, and hosting by verified organisations.",
        "Petitions: citizens propose changes and gather supporters toward a target, and administrators publish an official response.",
        "Community contributions: verified residents and organisations share completed work such as clean-up drives, volunteering and donation drives with a required photograph. Posting is limited per week, and others can applaud a post.",
    ]),
    ("h2", "Sustainable Funding"),
    ("p", "Two separate, clearly labelled mechanisms fund the service without tracking users. Sponsored posts let "
          "administrators feature vetted NGOs with a donation link, and each post must include an image or video. A "
          "self-serve advertising marketplace lets businesses and NGOs create campaigns that move through draft, "
          "review, approval and payment states. Ads are rotated in proportion to budget and matched only to the "
          "category of report being viewed, never to the person. Views and clicks are counted anonymously at most once "
          "per network per day, and advertisers are charged a flat rate of INR 100 per thousand views. Advertisers see "
          "their views, clicks, click-through rate and spend, and a campaign cannot be submitted for review without media."),

    ("h1", "Security and Privacy"),
    ("p", "Security rules are enforced in the database through PostgreSQL row-level security policies [5] and "
          "column-level grants, so that they cannot be bypassed by calling the API directly. Table IV lists the main controls."),
    ("table", "Security and Privacy Controls",
        ["Threat", "Control"],
        [
            ["Exposure of reporter identity", "Author column not readable by the public; anonymous reports show no name"],
            ["Leak of phone, address or ID digits", "Stored in private tables; encrypted with pgcrypto; decryption logged"],
            ["Privilege escalation", "Only an existing admin can grant the admin role; no self role changes"],
            ["Automated abuse", "CAPTCHA, rate limits per user, device and day"],
            ["Unreviewed content", "Community flagging with automatic hiding and admin moderation"],
            ["Secret exposure", "API keys stored in the database vault, never in the client"],
            ["Keeping data too long", "Nightly purge of expired personal and log data"],
            ["Unaccountable staff actions", "Every admin change and data reveal logged to an append-only audit table"],
            ["Spreadsheet formula injection", "Exported CSV cells are neutralised before download"],
        ]),
    ("p", "Identity verification accepts a masked Aadhaar document in which only the last four digits are visible, as "
          "issued by UIDAI [9]; the full twelve-digit number is never requested, in line with the Aadhaar Act [10]. "
          "The last four digits are encrypted at rest, and each time an administrator reveals them or a guest's contact "
          "details, the action is written to an append-only audit log, which also records every administrator change "
          "to content tables. Users can download a copy of their data or delete their account; deleting an account "
          "keeps the public reports but makes them anonymous. Retention limits are enforced by a nightly job: guest "
          "reporter contact details are deleted after 12 months, notifications after six months, application error "
          "logs after 30 days, missing child sightings 30 days after the alert closes, and raw ad view records after "
          "one day. Client error reports are scrubbed of email addresses, phone numbers and tokens before they are "
          "stored. These measures follow the principles of the Digital Personal Data Protection Act, 2023 [11], and "
          "the published privacy policy names a grievance officer. Public report data, excluding personal details, "
          "is available as open data in CSV and JSON formats."),

    ("h1", "Administration and Scalability"),
    ("p", "The administrator console is organised into five groups. Work covers the ticket queue, a reports view with "
          "quick replies, support threads, verification review, moderation of flagged content and reusable standard "
          "replies. Safety covers missing child alerts and city alerts. Content covers events, directories, NGOs and "
          "sponsored posts, missions, bus routes and petitions. Money covers advertisement review, marking payments and "
          "verifying advertisers. Oversight covers user management, insights, the activity log and the application "
          "error log."),
    ("p", "The dashboard can be filtered by period, area and category, exported as CSV or printed as PDF, and compares "
          "each figure with the previous period. It shows open, overdue and unclaimed tickets, time to first response, "
          "days to fix, on-time rate, citizen satisfaction, report hotspots by area, the hours of the day when reports "
          "arrive, actions per administrator, usage of free-plan limits, and advertisements awaiting payment or about "
          "to end. Insights chart reported against resolved reports over 12 weeks and by category."),
    ("p", "User management was designed for a large user base. The user list is paginated on the server and returns a "
          "total count, and name and organisation search uses trigram indexes so that partial-text search stays fast as "
          "the number of users grows. Administrators can change a user's role, account type or blocked status. An "
          "administrator cannot change their own role, and only an existing administrator can make another user an "
          "administrator."),

    ("h1", "Evaluation"),
    ("p", "Because CivicPulse has not yet been deployed to residents, it was evaluated at the system level on five "
          "dimensions: functional correctness, accuracy of the AI features, front-end performance and accessibility, "
          "back-end response time, and usability against established heuristics. It is also compared with existing platforms."),
    ("h2", "Functional Testing"),
    ("p", "The client passes TypeScript type checking and 41 automated unit tests, which run with a production build "
          "and a dependency audit on every push. Database rules were validated with rolled-back transactions that "
          "impersonate specific users, confirming that each policy allows and denies exactly what it should without "
          "leaving test data behind. End-to-end functional testing covered creating, editing and deleting records in "
          "each administrative module, guest and signed-in reporting, nested replies, identity verification from "
          "submission to approval, the dispute and reopen flow, and all three AI features against the live model. Every "
          "administrative change was independently confirmed in the server-side audit log. Testing surfaced and fixed "
          "several defects, including a disabled anonymous sign-in setting that silently blocked guest reports, a "
          "timeout that made department routing fail intermittently, a missing read permission on the email alert "
          "preference, and an unpaginated user list that would not have scaled beyond the first page of users."),
    ("h2", "AI Accuracy"),
    ("p", "Department routing was evaluated on a hand-labelled set of 40 Hyderabad reports, five for each of the "
          "eight departments the console offers, using the same prompt, output schema and model as the application. "
          "Of the first 18 cases evaluated, 17 were routed correctly (94.4%); the one miss sent an overflowing open drain to the water board rather than municipal sanitation, a genuinely ambiguous case in Hyderabad where the water board also handles sewerage. Evaluation of the remaining cases and of the duplicate check is limited by the free tier's quota of 20 requests per day [21] and is ongoing."),
    ("h2", "Performance and Accessibility"),
    ("p", "The production build was audited with Lighthouse 12 [19] under its default simulated mobile (slow 4G, "
          "throttled CPU) and desktop profiles. The first audit exposed two problems: a missing-child alert that loaded "
          "about two seconds after the rest of the home page pushed all content down by roughly 480 pixels, giving a "
          "cumulative layout shift of 0.655, and demonstration photographs served at full resolution. The home page was "
          "changed to wait for both alert queries before laying out, the sample photographs were resized, and a "
          "robots.txt file was added. Table V shows the scores before and after these fixes. Accessibility and best "
          "practices scored 100 in both profiles. The remaining mobile gap is the largest contentful paint of 5.2 s "
          "under simulated slow 4G, which is dominated by downloading the JavaScript bundle."),
    ("table", "Lighthouse Results (Before and After Fixes)",
        ["Metric", "Mobile", "Desktop"],
        [
            ["Performance score", "54 to 76", "84 to 99"],
            ["Accessibility score", "100", "100"],
            ["Best practices score", "100", "100"],
            ["SEO score", "90 to 100", "90 to 100"],
            ["Cumulative layout shift", "0.655 to 0", "0.307 to 0.025"],
            ["Total blocking time", "0 ms", "0 ms"],
            ["Page weight", "3.7 MB to 1.1 MB", "3.7 MB to 1.1 MB"],
        ]),
    ("h2", "Back-End Response Time"),
    ("p", "Loading the home page issues 13 requests to the managed backend (7 table reads and 6 function calls). "
          "Measured from the browser, their median round-trip time was 233 ms and the 90th percentile 252 ms, including "
          "network transit, row-level security checks and authentication. Calls to the language model took a median "
          "of 2.4 s (90th percentile 3.6 s), which is why every AI feature shows a progress state and never blocks "
          "submission. These figures were measured on the pilot-sized dataset; behaviour under city-scale data volumes "
          "has not yet been load-tested."),
    ("h2", "Heuristic Usability Evaluation"),
    ("p", "In place of a user study, the interface was reviewed against Nielsen's ten usability heuristics [20]. Table "
          "VI lists how each heuristic is addressed and the main weakness found."),
    ("table", "Heuristic Evaluation Against Nielsen's Heuristics",
        ["Heuristic", "How CivicPulse addresses it", "Weakness found"],
        [
            ["Visibility of system status", "Status stepper, public timeline, overdue marks, progress states for AI", "Queued offline reports are not listed"],
            ["Match with the real world", "Plain-language categories, 23 Indian languages, local helplines", "5 translations are drafts"],
            ["User control and freedom", "Edit or delete while pending, dismissible AI suggestions, dispute a fix", "No edits once triage starts"],
            ["Consistency and standards", "Shared card, badge and button styles across feed, profile and admin", "None significant"],
            ["Error prevention", "Duplicate warnings, 1.2 s hold on SOS, required proof photo", "None significant"],
            ["Recognition rather than recall", "Category icons, quick-detail chips, recent reports on the home page", "None significant"],
            ["Flexibility and efficiency", "Grid, list and map views; filters; standard staff replies", "AI answers take 2 to 4 s"],
            ["Aesthetic and minimalist design", "One primary action per screen, progressive disclosure", "Home page is long on phones"],
            ["Help users recover from errors", "Technical errors rewritten into plain language", "AI quota errors were unclear (fixed)"],
            ["Help and documentation", "Help and feedback threads, field hints", "No guided first-use tour"],
        ]),
    ("h2", "Comparison With Existing Platforms"),
    ("p", "Table VII compares CivicPulse with FixMyStreet and SeeClickFix on accountability and assistance features, "
          "based on their public documentation."),
    ("table", "Feature Comparison",
        ["Feature", "CivicPulse", "FixMyStreet", "SeeClickFix"],
        [
            ["Map-based reporting with photo", "Yes", "Yes", "Yes"],
            ["Reporter feedback after closure", "Confirm or dispute within 7 days", "Survey after 4 weeks", "Reopen within about 7 days"],
            ["After photo required to close", "Yes, enforced in database", "No", "No"],
            ["AI-assisted reporting and routing", "Yes", "Not documented", "Not documented"],
            ["Safety alerts and SOS", "Yes", "Not documented", "Not documented"],
        ]),

    ("h1", "Challenges Faced"),
    ("p", "Table VIII summarises the main technical, practical and legal challenges met during development and how each was resolved."),
    ("table", "Key Challenges and How They Were Resolved",
        ["Challenge", "Resolution"],
        [
            ["Security logic scattered across server code in the prototype", "Rebuilt with rules inside the database (row-level security, triggers)"],
            ["Browser storage quota overflow from client caching", "Service-worker caching with limits and an IndexedDB outbox"],
            ["No budget for paid services", "Free tiers only; per-user and global daily AI limits; SMS deferred"],
            ["Free AI quota of 20 requests a day", "Clear message when used up; AI never required to submit"],
            ["Home page jumping as alerts loaded", "Layout waits for alert queries; layout shift reduced from 0.655 to 0"],
            ["AI model retired during development", "Model name moved to a vault setting, switched without a release"],
            ["Unpredictable free-text AI answers", "Structured JSON output checked against a fixed schema"],
            ["Guest reporting silently failing", "Traced to a disabled anonymous sign-in setting; fixed and tested"],
            ["Intermittent AI routing timeouts", "Request timeout raised and results cached by text hash"],
            ["Admin user list not scaling", "Server-side pagination with trigram search indexes"],
            ["Legal limits on identity data", "Only masked Aadhaar accepted; last four digits encrypted; reveals audited"],
            ["No official GHMC interface", "One-tap hand-off to GHMC WhatsApp and grievance portal"],
            ["Translating into 23 languages", "20 translated; 3 fall back to English; 5 marked as drafts for review"],
        ]),

    ("h1", "Case Study Outcomes and Discussion"),
    ("p", "Table IX assesses each objective against what was built and how it was verified. The assessment is based on "
          "the implemented system and functional testing; measuring the effect on real resolution times requires the "
          "pilot described under Future Scope."),
    ("table", "Objectives Versus Outcomes",
        ["Objective", "Outcome and evidence"],
        [
            ["O1 Transparency", "Met: public reference, timeline, promised date, overdue marking and department scorecard"],
            ["O2 Proof of resolution", "Met: database trigger blocks resolving without a photo; dispute reopens the ticket (tested end to end)"],
            ["O3 Low effort, inclusion", "Met: guest reporting, AI auto-fill, offline outbox, 23 languages (20 translated)"],
            ["O4 Less noise", "Met: 250 m suggestions, 50 m submission guard, AI semantic check and admin merge"],
            ["O5 Privacy and safety", "Met: 124 row-level policies, encryption, audit log, retention purge; verified by impersonation tests"],
            ["O6 Sustainability", "Partly met: runs on free tiers with AI limits; sponsorship and ads are built but not yet proven with real revenue"],
        ]),
    ("h2", "Lessons Learned"),
    ("bullets", [
        "Enforce rules where the data lives: moving the proof-photo, rate-limit and privacy rules into the database made them impossible to bypass and simplified the client.",
        "Treat AI as a suggestion, not a decision: structured JSON output, a fixed schema, a human confirmation step and daily limits made a free-tier model dependable enough for everyday use.",
        "Keep the model configurable: when a model was retired during development, changing one vault setting moved the platform to a newer model without a release.",
        "Design for scale early: the admin user list was redesigned with server-side pagination and trigram search after testing showed it would not scale beyond the first page.",
        "Test the unhappy paths: disputes, guests, timeouts and missing permissions exposed more defects than the normal reporting flow did.",
    ]),

    ("h1", "Limitations"),
    ("bullets", [
        "No field evaluation yet: the system has been tested functionally but not with real residents and ward offices, so its effect on resolution time and satisfaction is not yet measured.",
        "No direct government integration: CivicPulse is a community platform and hands reports to GHMC only through its public WhatsApp and grievance portal, so a fix still depends on the department acting on that hand-off.",
        "Dependence on free tiers: the free Gemini tier allows only 20 requests per model per day for the whole project [21], so AI assistance is used up quickly under real load; the app now tells users plainly when this happens, but a paid plan or a lighter model is needed for a city-wide launch.",
        "Incomplete languages: three of the 23 languages (Bodo, Manipuri and Santali) have no translation yet, and five more need review by native speakers.",
        "No SMS notifications: SMS in India has no free tier, so residents without data access cannot receive updates.",
        "Single city: categories, helplines, bus data and alerts are specific to Hyderabad.",
        "Some administrative lists other than users still use fixed limits rather than server-side pagination.",
    ]),

    ("h1", "Future Scope"),
    ("bullets", [
        "Pilot study: deploy with residents and one or more GHMC wards and measure time to fix, the share of fixes confirmed by reporters and user satisfaction.",
        "Official integration: an Open311-compatible interface [1] so that reports reach municipal systems directly, with status updates flowing back automatically.",
        "Native mobile apps for Android and iOS, with background location and richer push notifications.",
        "More AI assistance: on-demand translation of user-written text, voice reporting in local languages, image checks for before and after photos, and predictive maps of likely problem areas from past reports.",
        "SMS and IVR channels for residents without smartphones, once a funding source is in place.",
        "Multi-city support, with each city configuring its own departments, categories, helplines and alerts.",
        "Open data dashboards for researchers, journalists and ward committees, built on the existing CSV and JSON export.",
        "Complete the remaining translations with native speakers and add screen-reader testing for accessibility.",
    ]),

    ("h1", "Conclusion"),
    ("p", "CivicPulse shows that a city grievance platform can be both easy to use and accountable. Mandatory before and "
          "after photographs, reporter confirmation and public timelines make it hard to close a ticket without fixing "
          "the problem. This case study followed the full journey of a report, from a photograph taken on a street to a fix confirmed by the person who reported it, and assessed how each design objective was met. Carefully bounded AI assistance reduces the effort of reporting and triage without taking "
          "decisions away from people, and enforcing privacy and access rules in the database keeps personal data safe "
          "even if the client is bypassed. The platform is ready for a pilot deployment in Hyderabad."),
]

ACK = ("The authors thank the open-source communities behind React, PostgreSQL, Supabase, Leaflet and OpenStreetMap, "
       "whose work made this project possible. An AI coding assistant was used during software development and in "
       "preparing this manuscript; all content was reviewed and verified by the authors, who take full responsibility for it.")

REFS = [
    'Open311, "GeoReport v2 specification." [Online]. Available: https://www.open311.org/',
    'mySociety, "FixMyStreet." [Online]. Available: https://www.fixmystreet.com/',
    'SeeClickFix. [Online]. Available: https://seeclickfix.com/',
    'Ministry of Housing and Urban Affairs, Government of India, "Swachhata-MoHUA mobile application."',
    'The PostgreSQL Global Development Group, "Row security policies," PostgreSQL documentation. [Online]. Available: https://www.postgresql.org/docs/current/ddl-rowsecurity.html',
    'Supabase, "Supabase documentation." [Online]. Available: https://supabase.com/docs',
    'Google, "Gemini API documentation." [Online]. Available: https://ai.google.dev/gemini-api/docs',
    'OpenStreetMap Foundation, "Nominatim." [Online]. Available: https://nominatim.org/',
    'Unique Identification Authority of India, "Masked Aadhaar." [Online]. Available: https://uidai.gov.in/',
    'The Aadhaar (Targeted Delivery of Financial and Other Subsidies, Benefits and Services) Act, 2016, Government of India.',
    'The Digital Personal Data Protection Act, 2023, Government of India.',
    'Cloudflare, "Turnstile documentation." [Online]. Available: https://developers.cloudflare.com/turnstile/',
    'D. Linders, "From e-government to we-government: Defining a typology for citizen coproduction in the age of social media," Government Information Quarterly, vol. 29, no. 4, pp. 446-454, 2012.',
    'M. F. Goodchild, "Citizens as sensors: The world of volunteered geography," GeoJournal, vol. 69, no. 4, pp. 211-221, 2007.',
    'T. Nam and T. A. Pardo, "Conceptualizing smart city with dimensions of technology, people, and institutions," in Proc. 12th Annu. Int. Digital Government Research Conf. (dg.o), 2011, pp. 282-291.',
    'S. F. King and P. Brown, "Fix my street or else: Using the internet to voice local public service concerns," in Proc. 1st Int. Conf. Theory and Practice of Electronic Governance (ICEGOV), 2007.',
    'mySociety, "FixMyStreet glossary: questionnaire." [Online]. Available: https://fixmystreet.org/glossary/',
    'CivicPlus, "How do I reopen an issue that has been closed?" SeeClickFix help centre. [Online]. Available: https://www.seeclickfixusers.civicplus.help/hc/en-us/articles/360043575173-How-do-I-reopen-an-Issue-that-has-been-closed',
    'Google, "Lighthouse overview," Chrome for Developers. [Online]. Available: https://developer.chrome.com/docs/lighthouse/overview',
    'J. Nielsen, "Enhancing the explanatory power of usability heuristics," in Proc. SIGCHI Conf. Human Factors in Computing Systems (CHI), 1994, pp. 152-158.',
    'Google, "Gemini API rate limits." [Online]. Available: https://ai.google.dev/gemini-api/docs/rate-limits',
]


# ---------------------------------------------------------------- figures
def _font(size, bold=False):
    return ImageFont.truetype(FONT_B if bold else FONT, size)


def _box(d, xy, text, fill="#ffffff", outline="#1f2937", bold=False, size=26):
    d.rounded_rectangle(xy, radius=14, fill=fill, outline=outline, width=3)
    f = _font(size, bold)
    lines = text.split("\n")
    lh = size + 6
    x0, y0, x1, y1 = xy
    y = (y0 + y1) / 2 - lh * len(lines) / 2
    for ln in lines:
        w = d.textlength(ln, font=f)
        d.text(((x0 + x1) / 2 - w / 2, y), ln, font=f, fill="#111827")
        y += lh


def _arrow(d, a, b, label=None):
    d.line([a, b], fill="#374151", width=4)
    import math
    ang = math.atan2(b[1] - a[1], b[0] - a[0])
    for s in (-0.45, 0.45):
        d.line([b, (b[0] - 18 * math.cos(ang + s), b[1] - 18 * math.sin(ang + s))], fill="#374151", width=4)
    if label:
        f = _font(22)
        mx, my = (a[0] + b[0]) / 2, (a[1] + b[1]) / 2
        w = d.textlength(label, font=f)
        d.rectangle([mx - w / 2 - 4, my - 14, mx + w / 2 + 4, my + 14], fill="#ffffff")
        d.text((mx - w / 2, my - 12), label, font=f, fill="#374151")


def figure_architecture():
    img = Image.new("RGB", (1400, 1000), "white")
    d = ImageDraw.Draw(img)
    # Client tier
    d.rounded_rectangle([40, 30, 1360, 260], radius=18, outline="#2563eb", width=3, fill="#eff6ff")
    d.text((60, 40), "Client: React PWA (TypeScript, Tailwind, Vite)", font=_font(30, True), fill="#1e3a8a")
    _box(d, (70, 100, 420, 230), "Citizen app\nreport, feed, SOS")
    _box(d, (525, 100, 875, 230), "Admin console\nqueue, moderation")
    _box(d, (980, 100, 1330, 230), "Service worker\noffline outbox")
    # Backend tier
    d.rounded_rectangle([40, 340, 1360, 700], radius=18, outline="#059669", width=3, fill="#ecfdf5")
    d.text((60, 350), "Supabase backend", font=_font(30, True), fill="#065f46")
    _box(d, (70, 410, 400, 520), "Auth\n(incl. anonymous)")
    _box(d, (440, 410, 960, 520), "REST API (PostgREST)\nRPC functions")
    _box(d, (1000, 410, 1330, 520), "Storage\nphotos, videos")
    _box(d, (70, 560, 1330, 680), "PostgreSQL: RLS policies, triggers, pgcrypto, Vault secrets, pg_cron, http / pg_net", bold=True, size=25)
    # External tier
    d.rounded_rectangle([40, 780, 1360, 970], radius=18, outline="#b45309", width=3, fill="#fffbeb")
    d.text((1060, 790), "External services", font=_font(30, True), fill="#78350f")
    _box(d, (60, 840, 300, 950), "Gemini API\n(AI)")
    _box(d, (320, 840, 560, 950), "Resend\n(email)")
    _box(d, (580, 840, 820, 950), "OSM Nominatim\nOverpass")
    _box(d, (840, 840, 1080, 950), "Open-Meteo\n(weather, AQI)")
    _box(d, (1100, 840, 1340, 950), "Turnstile\n(CAPTCHA)")
    _arrow(d, (700, 260), (700, 410), "HTTPS + JWT")
    _arrow(d, (180, 680), (180, 840), "sync")
    _arrow(d, (440, 680), (440, 840), "async")
    img.save(OUT / "fig_architecture.png", dpi=(300, 300))


def _panel(d, xy, title, outline, fill, title_fill):
    d.rounded_rectangle(xy, radius=20, outline=outline, width=4, fill=fill)
    d.text((xy[0] + 24, xy[1] + 14), title, font=_font(34, True), fill=title_fill)


def _chip(d, xy, text, fill="#ffffff", size=24, bold=False):
    _box(d, xy, text, fill=fill, size=size, bold=bold)


def figure_architecture_detailed():
    W, H = 2400, 1900
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)

    # 1. Users
    _panel(d, (40, 30, W - 40, 190), "Users", "#6b7280", "#f9fafb", "#374151")
    for i, u in enumerate(["Citizen / guest", "Verified organisation", "Administrator", "Advertiser"]):
        x = 360 + i * 500
        _chip(d, (x, 80, x + 420, 170), u, fill="#f3f4f6", size=28, bold=True)

    # 2. Client
    _panel(d, (40, 230, W - 40, 760), "Client: React 19 progressive web app (TypeScript, Tailwind CSS, Vite)", "#2563eb", "#eff6ff", "#1e3a8a")
    mods = [
        ("Feed and map", "Report + AI assist", "Report detail\ntimeline, replies"),
        ("Events, petitions\nmissions", "Utilities, buses\nsupport", "Profile\nverification"),
        ("SOS and\nchild alerts", "Admin console\n(5 groups)", "Advertiser\ncampaigns"),
    ]
    for r, row in enumerate(mods):
        for c, m in enumerate(row):
            x, y = 80 + c * 420, 300 + r * 145
            _chip(d, (x, y, x + 390, y + 120), m, size=25)
    d.text((1370, 300), "Cross-cutting", font=_font(28, True), fill="#1e3a8a")
    cross = ["TanStack Query cache", "React Router", "i18n: 23 languages, RTL", "Leaflet maps, exifr",
             "Workbox service worker", "IndexedDB offline outbox", "Web Push subscription", "Turnstile widget"]
    for i, t in enumerate(cross):
        x = 1370 + (i % 2) * 480
        y = 350 + (i // 2) * 98
        _chip(d, (x, y, x + 450, y + 80), t, fill="#dbeafe", size=24)

    # 3. Supabase platform
    _panel(d, (40, 830, W - 40, 1520), "Supabase platform", "#059669", "#ecfdf5", "#065f46")
    svc = [("Auth\nemail, guest, Google", "#d1fae5"), ("REST API (PostgREST)\ntables + RPC", "#d1fae5"),
           ("Storage buckets\nphotos, videos, ads,\nalerts, documents", "#d1fae5")]
    for i, (t, f) in enumerate(svc):
        x = 80 + i * 760
        _chip(d, (x, 890, x + 700, 1030), t, fill=f, size=26, bold=True)
    d.rounded_rectangle((80, 1070, W - 80, 1490), radius=16, outline="#047857", width=3, fill="#ffffff")
    d.text((110, 1085), "PostgreSQL: 44 tables, 109 functions, 53 triggers, 124 row-level security policies", font=_font(28, True), fill="#065f46")
    db = [
        "Row-level security\ncolumn-level grants", "Triggers: lifecycle,\ndedupe, rate limits", "Triggers: notify,\naudit, points",
        "RPC: AI assist,\nadmin, claims", "pgcrypto encryption\n+ Vault secrets",
        "http (sync)\nAI calls", "pg_net (async)\nemail, push", "pg_cron\nnightly purge", "pg_trgm\nsearch indexes",
    ]
    for i, t in enumerate(db):
        col, row = i % 5, i // 5
        x = 110 + col * 440
        y = 1150 + row * 165
        _chip(d, (x, y, x + 410, y + 140), t, fill="#ecfdf5", size=25)

    # 4. External services
    d.rounded_rectangle((40, 1590, W - 40, 1870), radius=20, outline="#b45309", width=4, fill="#fffbeb")
    d.text((64, 1812), "External services", font=_font(34, True), fill="#78350f")
    ext = ["Google Gemini\n(AI)", "Resend\n(email)", "Browser push\nservices", "OpenStreetMap\nNominatim, Overpass",
           "Open-Meteo\nweather, AQI", "Cloudflare\nTurnstile"]
    for i, t in enumerate(ext):
        x = 80 + i * 380
        _chip(d, (x, 1640, x + 350, 1800), t, fill="#fef3c7", size=26)

    # Flows
    _arrow(d, (1200, 190), (1200, 300), "HTTPS")
    _arrow(d, (430, 760), (430, 890), "sign in")
    _arrow(d, (1190, 760), (1190, 890), "JWT-authenticated REST / RPC")
    _arrow(d, (1950, 760), (1950, 890), "uploads")
    # From the database's HTTP extensions to AI, email and push
    _arrow(d, (315, 1455), (255, 1640), "sync")
    _arrow(d, (755, 1455), (635, 1640), "async")
    _arrow(d, (815, 1455), (1015, 1640))
    # The browser calls map, weather and CAPTCHA services directly, not through the backend
    gx = W - 55
    d.line([(W - 40, 700), (gx, 700), (gx, 1612)], fill="#6b7280", width=4)
    d.line([(1395, 1615), (gx, 1612)], fill="#6b7280", width=4)
    for tx in (1395, 1775, 2155):
        _arrow(d, (tx, 1612), (tx, 1640))
    f = _font(24)
    lbl = "direct from browser"
    w = d.textlength(lbl, font=f)
    d.rectangle([1700 - 6, 1597, 1700 + w + 6, 1627], fill="#ffffff")
    d.text((1700, 1598), lbl, font=f, fill="#374151")
    img.save(OUT / "fig_architecture.png", dpi=(300, 300))


def figure_lifecycle():
    img = Image.new("RGB", (1400, 570), "white")
    d = ImageDraw.Draw(img)
    _box(d, (40, 80, 300, 200), "Pending\n(reported)", fill="#fef3c7")
    _box(d, (400, 80, 660, 200), "In progress\n(claimed, dated)", fill="#dbeafe")
    _box(d, (760, 80, 1020, 200), "Resolved\n(proof photo)", fill="#dcfce7")
    _box(d, (1100, 30, 1370, 140), "Accepted\nby reporter", fill="#bbf7d0")
    _box(d, (1100, 170, 1370, 280), "Auto-confirmed\nafter 7 days", fill="#e5e7eb")
    _box(d, (760, 420, 1020, 540), "Disputed\n(reopened)", fill="#fee2e2")
    _box(d, (220, 420, 560, 540), "Closed, not fixed\n(stated reason)", fill="#e5e7eb")
    _arrow(d, (300, 140), (400, 140))
    _arrow(d, (660, 140), (760, 140), "photo")
    _arrow(d, (1020, 120), (1100, 85))
    _arrow(d, (1020, 160), (1100, 225))
    _arrow(d, (890, 200), (890, 420), "reporter disputes")
    _arrow(d, (760, 480), (560, 200), "back to work")
    _arrow(d, (170, 200), (330, 420), "reason")
    _arrow(d, (530, 200), (430, 420))
    img.save(OUT / "fig_lifecycle.png", dpi=(300, 300))


# ---------------------------------------------------------------- LaTeX
def tex_escape(s):
    rep = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_", "{": r"\{", "}": r"\}",
           "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}
    out = "".join(rep.get(c, c) for c in s)
    # "quoted" -> ``quoted''
    parts = out.split('"')
    return "".join(p if i % 2 == 0 else "``" + p + "''" for i, p in enumerate(parts))


def tex_prose(s):
    txt = tex_escape(s)
    txt = txt.replace("Fig. 1", r"Fig.~\ref{fig:arch}").replace("Fig. 2", r"Fig.~\ref{fig:life}")
    romans = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6, "VII": 7, "VIII": 8, "IX": 9, "X": 10}
    txt = re.sub(r"Table (X|IX|VIII|VII|VI|IV|V|I{1,3})\b", lambda m: r"Table~\ref{tab:%d}" % romans[m.group(1)], txt)
    for i in range(len(REFS), 0, -1):
        txt = txt.replace(f"[{i}]", r"\cite{b" + str(i) + "}")
    return txt


def build_tex():
    L = [r"\documentclass[conference]{IEEEtran}", r"\IEEEoverridecommandlockouts",
         r"\usepackage{cite}", r"\usepackage{graphicx}", r"\usepackage{textcomp}", r"\usepackage{xcolor}",
         r"\usepackage{array}", r"\usepackage{url}", r"\begin{document}",
         r"\title{" + tex_escape(TITLE) + "}", r"\author{"]
    blocks = []
    for i, (n, dept, org, city, mail) in enumerate(AUTHORS):
        ord_ = ["1st", "2nd", "3rd", "4th", "5th", "6th"][i]
        blocks.append(r"\IEEEauthorblockN{" + tex_escape(n) + "}\n" + r"\IEEEauthorblockA{\textit{" + tex_escape(dept) + r"} \\" +
                      "\n" + r"\textit{" + tex_escape(org) + r"}\\" + "\n" + tex_escape(city) + r" \\" + "\n" + tex_escape(mail) + "}")
    L.append("\n\\and\n".join(blocks))
    L += ["}", r"\maketitle", r"\begin{abstract}", tex_escape(ABSTRACT), r"\end{abstract}",
          r"\begin{IEEEkeywords}", tex_escape(INDEX_TERMS), r"\end{IEEEkeywords}"]
    fig_n = 0
    tab_n = 0
    for b in BODY:
        k = b[0]
        if k == "h1":
            L.append(r"\section{" + tex_escape(b[1]) + "}")
        elif k == "h2":
            L.append(r"\subsection{" + tex_escape(b[1]) + "}")
        elif k == "p":
            L.append(tex_prose(b[1]) + "\n")
        elif k == "bullets":
            L.append(r"\begin{itemize}")
            L += [r"\item " + tex_prose(x) for x in b[1]]
            L.append(r"\end{itemize}")
        elif k == "table":
            tab_n += 1
            _, cap, head, rows = b
            cols = len(head)
            spec = "|" + "|".join([r">{\raggedright\arraybackslash}p{%.2f\columnwidth}" % (0.9 / cols if cols > 2 else (0.3 if j == 0 else 0.6)) for j in range(cols)]) + "|"
            L += [r"\begin{table}[htbp]", r"\caption{" + tex_escape(cap) + "}", r"\label{tab:%d}" % tab_n, r"\begin{center}",
                  r"\footnotesize", r"\begin{tabular}{" + spec + "}", r"\hline",
                  " & ".join(r"\textbf{" + tex_escape(h) + "}" for h in head) + r" \\", r"\hline"]
            L += [" & ".join(tex_escape(c) for c in r) + r" \\ \hline" for r in rows]
            L += [r"\end{tabular}", r"\end{center}", r"\end{table}"]
        elif k in ("fig", "fig_wide"):
            fig_n += 1
            lab = "fig:arch" if fig_n == 1 else "fig:life"
            env, width = ("figure*", r"\textwidth") if k == "fig_wide" else ("figure", r"\columnwidth")
            L += [r"\begin{" + env + "}[htbp]", r"\centerline{\includegraphics[width=" + width + "]{" + (b[1].replace(".png", ".pdf") if k == "fig_wide" else b[1]) + "}}",
                  r"\caption{" + tex_escape(b[2]) + "}", r"\label{" + lab + "}", r"\end{" + env + "}"]
    L += [r"\section*{Acknowledgment}", tex_escape(ACK), r"\begin{thebibliography}{00}"]
    L += [r"\bibitem{b%d} " % (i + 1) + tex_escape(r).replace("Available: ", "Available: \\url{").replace("https://", "https://") +
          ("}" if "Available: " in r else "") for i, r in enumerate(REFS)]
    L += [r"\end{thebibliography}", r"\end{document}"]
    tex = "\n".join(L)
    # tex_escape turned URL underscores/percent into escaped forms; \url wants them raw.
    import re
    tex = re.sub(r"\\url\{([^}]*)\}", lambda m: r"\url{" + m.group(1).replace(r"\_", "_") + "}", tex)
    (OUT / "civicpulse_ieee.tex").write_text(tex, encoding="utf-8")


# ---------------------------------------------------------------- Word
def _set_cols(section, n):
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), "360")


def _run(p, text, size=10, bold=False, italic=False, smallcaps=False):
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    r.font.size = Pt(size)
    r.bold, r.italic = bold, italic
    r.font.small_caps = smallcaps
    return r


def _para(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=4, indent=True):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = 1.0
    if indent:
        pf.first_line_indent = Cm(0.35)
    return p


ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII", "XIII", "XIV", "XV", "XVI", "XVII", "XVIII", "XIX", "XX"]


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(1.0)
    sec.left_margin = sec.right_margin = Inches(0.625)

    p = _para(doc, WD_ALIGN_PARAGRAPH.CENTER, after=12, indent=False)
    _run(p, TITLE, size=24)
    t = doc.add_table(rows=1, cols=len(AUTHORS))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, (n, dept, org, city, mail) in zip(t.rows[0].cells, AUTHORS):
        cp = c.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(cp, n + "\n", size=11)
        _run(cp, dept + "\n" + org + "\n", size=10, italic=True)
        _run(cp, city + "\n" + mail, size=10)

    body = doc.add_section(WD_SECTION.CONTINUOUS)
    _set_cols(body, 2)

    p = _para(doc, before=10)
    _run(p, "Abstract", size=9, bold=True, italic=True)
    _run(p, "\u2014" + ABSTRACT, size=9, bold=True)
    p = _para(doc, after=8)
    _run(p, "Index Terms", size=9, bold=True, italic=True)
    _run(p, "\u2014" + INDEX_TERMS, size=9, bold=True)

    h1 = h2 = tab_n = fig_n = 0
    for b in BODY:
        k = b[0]
        if k == "h1":
            h2 = 0
            p = _para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=8, after=4, indent=False)
            _run(p, f"{ROMAN[h1]}. {b[1]}", size=10, smallcaps=True)
            h1 += 1
        elif k == "h2":
            p = _para(doc, WD_ALIGN_PARAGRAPH.LEFT, before=4, after=2, indent=False)
            _run(p, f"{chr(65 + h2)}. {b[1]}", size=10, italic=True)
            h2 += 1
        elif k == "p":
            _run(_para(doc), b[1])
        elif k == "bullets":
            for x in b[1]:
                p = _para(doc, indent=False, after=2)
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.first_line_indent = Cm(-0.3)
                _run(p, "\u2022 " + x)
        elif k == "table":
            tab_n += 1
            _, cap, head, rows = b
            p = _para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=6, after=2, indent=False)
            _run(p, f"TABLE {ROMAN[tab_n - 1]}\n{cap.upper()}", size=8, smallcaps=True)
            tb = doc.add_table(rows=1 + len(rows), cols=len(head))
            tb.style = "Table Grid"
            tb.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, h in enumerate(head):
                cp = tb.rows[0].cells[j].paragraphs[0]
                _run(cp, h, size=8, bold=True)
            for i, row in enumerate(rows):
                for j, v in enumerate(row):
                    _run(tb.rows[i + 1].cells[j].paragraphs[0], v, size=8)
            _para(doc, after=2, indent=False)
        elif k == "fig_wide":
            fig_n += 1
            _set_cols(doc.add_section(WD_SECTION.CONTINUOUS), 1)
            p = _para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=6, after=0, indent=False)
            p.add_run().add_picture(str(OUT / b[1]), width=Inches(7.0))
            p = _para(doc, WD_ALIGN_PARAGRAPH.LEFT, after=6, indent=False)
            _run(p, f"Fig. {fig_n}. {b[2]}", size=8)
            _set_cols(doc.add_section(WD_SECTION.CONTINUOUS), 2)
        elif k == "fig":
            fig_n += 1
            p = _para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=6, after=0, indent=False)
            p.add_run().add_picture(str(OUT / b[1]), width=Inches(3.4))
            p = _para(doc, WD_ALIGN_PARAGRAPH.LEFT, after=6, indent=False)
            _run(p, f"Fig. {fig_n}. {b[2]}", size=8)

    p = _para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=8, after=4, indent=False)
    _run(p, "Acknowledgment", size=10, smallcaps=True)
    _run(_para(doc), ACK)
    p = _para(doc, WD_ALIGN_PARAGRAPH.CENTER, before=8, after=4, indent=False)
    _run(p, "References", size=10, smallcaps=True)
    for i, r in enumerate(REFS, 1):
        p = _para(doc, WD_ALIGN_PARAGRAPH.LEFT, after=1, indent=False)
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        _run(p, f"[{i}]  {r}", size=8)
    doc.save(OUT / "CivicPulse_IEEE.docx")


def renumber_refs():
    """IEEE numbers references in order of first citation; remap [n] in BODY and reorder REFS to match."""
    global BODY, REFS
    order = []
    def walk(x):
        if isinstance(x, str):
            for n in re.findall(r"\[(\d+)\]", x):
                if int(n) not in order:
                    order.append(int(n))
        elif isinstance(x, (list, tuple)):
            for y in x:
                walk(y)
    walk(BODY)
    order += [i for i in range(1, len(REFS) + 1) if i not in order]
    new_of = {old: new for new, old in enumerate(order, 1)}
    def sub(x):
        if isinstance(x, str):
            return re.sub(r"\[(\d+)\]", lambda m: f"[{new_of[int(m.group(1))]}]", x)
        if isinstance(x, tuple):
            return tuple(sub(y) for y in x)
        if isinstance(x, list):
            return [sub(y) for y in x]
        return x
    BODY = sub(BODY)
    REFS = [REFS[old - 1] for old in order]


if __name__ == "__main__":
    renumber_refs()
    import architecture_svg  # noqa: F401  (writes vector fig_architecture.svg/.pdf/.png)
    figure_lifecycle()
    build_tex()
    build_docx()
    print("Wrote:", *sorted(x.name for x in OUT.iterdir() if x.suffix in {".tex", ".docx", ".png"}))
